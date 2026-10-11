import openpyxl, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
wb = openpyxl.load_workbook(r"downloads\目押记录.xlsx", data_only=True)
for ws in wb.worksheets:
    print("=== sheet: %s  dims=%s  max_row=%d max_col=%d ===" % (ws.title, ws.dimensions, ws.max_row, ws.max_column))
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        cells = ["" if c is None else str(c) for c in row]
        if any(c.strip() for c in cells):
            print("  r%-3d | %s" % (i, " | ".join(cells)))
